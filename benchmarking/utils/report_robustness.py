"""Tables of the missing-modality results (test_robustness.py), one per experiment, printed
and saved to an .xlsx workbook.

    python utils/report_robustness.py [experiment substrings ...] [--logs_folder logs]
"""

import argparse
import glob
import json
import os
import re

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from tabulate import tabulate

RESULTS_FILENAME_PREFIX = "test_results_robustness_"

# longest first, so that e.g. "sensormae_base_vit" is matched before "vit"
MODEL_SUFFIXES = sorted(
    [
        "vit",
        "mae",
        "dino",
        "multimae",
        "sensormae_base_vit",
        "sensormae_base_mm",
        "sensormae_small_vit",
        "sensormae_tiny_vit",
    ],
    key=len,
    reverse=True,
)


def load_machine_config(logs_folder, machine_config_path="machine_config.yaml"):
    if os.path.exists(machine_config_path):
        with open(machine_config_path, "r") as f:
            machine_conf = yaml.safe_load(f)
        if "logs_folder" in machine_conf:
            return machine_conf["logs_folder"]
    return logs_folder


def split_experiment_and_model(dir_name):
    """'3_det_dsec_det_rgb+event+depth_sensormae_base_mm' ->
    ('3_det_dsec_det_rgb+event+depth', 'sensormae_base_mm'), (None, None) if unknown."""
    for suffix in MODEL_SUFFIXES:
        marker = "_" + suffix
        if dir_name.endswith(marker):
            return dir_name[: -len(marker)], suffix
    return None, None


def discover_experiment_groups(logs_folder):
    """{experiment_name: {model_suffix: dir_name}} of the folders in logs_folder."""
    groups = {}
    for entry in sorted(os.listdir(logs_folder)):
        full_path = os.path.join(logs_folder, entry)
        if not os.path.isdir(full_path):
            continue
        exp_name, model = split_experiment_and_model(entry)
        if exp_name is None:
            continue
        groups.setdefault(exp_name, {})[model] = entry
    return groups


def find_latest_results_file(logs_folder, exp_dir):
    pattern = os.path.join(
        logs_folder, exp_dir, "test_results", f"{RESULTS_FILENAME_PREFIX}*.json"
    )
    paths = glob.glob(pattern)
    if not paths:
        return None
    return max(paths, key=os.path.getmtime)


def load_group_results(logs_folder, model_to_dir):

    results = {}
    for model, dir_name in model_to_dir.items():
        path = find_latest_results_file(logs_folder, dir_name)
        if path is None:
            continue
        with open(path, "r") as f:
            results[model] = json.load(f)
    return results


def extract_metrics(metrics):
    """{metric: value in %}: mAP50:95 and mAP50 for detection, mIoU for segmentation."""
    if not metrics:
        return {}
    entry = metrics[0]
    if "test_map" in entry or "test_map_50" in entry:
        result = {}
        if "test_map" in entry:
            result["mAP50:95"] = entry["test_map"] * 100
        if "test_map_50" in entry:
            result["mAP50"] = entry["test_map_50"] * 100
        return result
    if "test_mIoU" in entry:
        return {"mIoU": entry["test_mIoU"] * 100}
    for key, value in entry.items():
        if isinstance(value, (int, float)):
            return {key: value * 100}
    return {}


def active_label(subset_data):
    if not subset_data["missing_modalities"]:
        return "all"
    return "+".join(sorted(subset_data["active_modalities"]))


def column_sort_key(subset_data):
    label = active_label(subset_data)
    if label == "all":
        return (1, 0, "")
    return (0, len(subset_data["active_modalities"]), label)


def build_group_table(model_results):
    """(headers, rows) of one experiment; missing values are "-"."""
    columns = {}
    for subset_results in model_results.values():
        for subset_data in subset_results.values():
            columns[column_sort_key(subset_data)] = active_label(subset_data)
    ordered_columns = [columns[key] for key in sorted(columns.keys())]

    headers = ["Model", "Metric"] + ordered_columns
    rows = []
    for model in MODEL_SUFFIXES:
        if model not in model_results:
            continue
        label_to_metrics = {
            active_label(subset_data): extract_metrics(subset_data["metrics"])
            for subset_data in model_results[model].values()
        }

        metric_labels = []
        for metrics in label_to_metrics.values():
            for metric_label in metrics:
                if metric_label not in metric_labels:
                    metric_labels.append(metric_label)

        for metric_label in metric_labels:
            row = [model, metric_label] + [
                label_to_metrics.get(col, {}).get(metric_label, "-")
                for col in ordered_columns
            ]
            rows.append(row)

    return headers, rows


def print_group_table(exp_name, headers, rows):
    print(f"\n=== {exp_name} ===")
    print(tabulate(rows, headers=headers, tablefmt="rounded_grid", floatfmt=".2f"))


def sanitize_sheet_name(name, used_names):
    sanitized = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
    candidate = sanitized
    suffix_idx = 2
    while candidate in used_names:
        suffix = f"_{suffix_idx}"
        candidate = sanitized[: 31 - len(suffix)] + suffix
        suffix_idx += 1
    used_names.add(candidate)
    return candidate


def write_group_sheet(wb, exp_name, headers, rows, used_names):
    ws = wb.create_sheet(title=sanitize_sheet_name(exp_name, used_names))

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row][2:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 10
    for i in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12


def write_excel(group_tables, output_path):
    """One sheet per experiment."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names = set()
    for exp_name, (headers, rows) in group_tables.items():
        write_group_sheet(wb, exp_name, headers, rows, used_names)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Report missing-modality robustness results, auto-grouped by "
            "experiment across all models found under logs_folder."
        )
    )
    parser.add_argument(
        "experiment_filters",
        type=str,
        nargs="*",
        help=(
            "Optional experiment name substrings to restrict which groups "
            "are reported (default: report every group found)."
        ),
    )
    parser.add_argument("--logs_folder", type=str, default="logs")
    parser.add_argument("--machine_config", type=str, default="machine_config.yaml")
    parser.add_argument(
        "--xlsx",
        type=str,
        default=None,
        help=(
            "Path to save an .xlsx report with one sheet per experiment "
            "group (default: <logs_folder>/robustness_report.xlsx)."
        ),
    )
    args = parser.parse_args()

    logs_folder = load_machine_config(args.logs_folder, args.machine_config)
    groups = discover_experiment_groups(logs_folder)

    if args.experiment_filters:
        groups = {
            exp_name: model_to_dir
            for exp_name, model_to_dir in groups.items()
            if any(f in exp_name for f in args.experiment_filters)
        }

    group_tables = {}
    for exp_name in sorted(groups.keys()):
        model_results = load_group_results(logs_folder, groups[exp_name])
        if not model_results:
            continue
        headers, rows = build_group_table(model_results)
        print_group_table(exp_name, headers, rows)
        group_tables[exp_name] = (headers, rows)

    if group_tables:
        xlsx_path = args.xlsx or os.path.join(logs_folder, "robustness_report.xlsx")
        write_excel(group_tables, xlsx_path)
        print(f"\nExcel report saved to: {xlsx_path}")
    else:
        print("No robustness results found; skipping Excel export.")


if __name__ == "__main__":
    main()
