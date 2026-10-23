"""Excel summary of the STF results (kitti_test_results_*.json from eval_stf.py).

  --mode full              easy / moderate / hard AP per weather condition for every
                           (modalities, backbone) pair; Car class and all classes.
  --mode scale-comparison  SensorMAE-ViT Base / Small / Tiny, AP averaged over the
                           difficulties.

    python utils/stf_results_to_excel.py --logs-dir logs --mode full
"""

import argparse
import glob
import json
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MODALITIES_ORDER = [
    "1_RGB",
    "1_SNIR",
    "1_DEPTH",
    "1_RADAR",
    "1_LIDAR",
    "2_RGB_SNIR",
    "2_RGB_DEPTH",
    "2_RGB_RADAR",
    "2_RGB_LIDAR",
    "2_SNIR_DEPTH",
    "2_SNIR_LIDAR",
    "3_RGB_SNIR_DEPTH",
    "3_RGB_SNIR_RADAR",
    "3_RGB_SNIR_LIDAR",
    "3_RGB_RADAR_LIDAR",
    "3_SNIR_DEPTH_RADAR",
    "3_RGB_SNIR_LIDAR_RADAR",
    "3_RGB_SNIR_DEPTH_RADAR",
    "3_RGB_SNIR_DEPTH_RADAR_LIDAR",
]

# sorted, '+'-joined modalities of the folder name -> row label
MODALITY_MAP = {
    "depth": "1_DEPTH",
    "lidar": "1_LIDAR",
    "radar": "1_RADAR",
    "rgb": "1_RGB",
    "snir": "1_SNIR",
    "depth+rgb": "2_RGB_DEPTH",
    "lidar+rgb": "2_RGB_LIDAR",
    "radar+rgb": "2_RGB_RADAR",
    "rgb+snir": "2_RGB_SNIR",
    "depth+snir": "2_SNIR_DEPTH",
    "lidar+snir": "2_SNIR_LIDAR",
    "depth+rgb+snir": "3_RGB_SNIR_DEPTH",
    "radar+rgb+snir": "3_RGB_SNIR_RADAR",
    "lidar+rgb+snir": "3_RGB_SNIR_LIDAR",
    "lidar+radar+rgb": "3_RGB_RADAR_LIDAR",
    "depth+radar+snir": "3_SNIR_DEPTH_RADAR",
    "lidar+radar+rgb+snir": "3_RGB_SNIR_LIDAR_RADAR",
    "depth+radar+rgb+snir": "3_RGB_SNIR_DEPTH_RADAR",
    "depth+lidar+radar+rgb+snir": "3_RGB_SNIR_DEPTH_RADAR_LIDAR",
}

# folder-name backbone suffix -> column label (longest suffixes are matched first)
MODEL_MAP = {
    "sensormae_base_vit": "Our-ViT",
    "sensormae_base_mm": "Our-MM",
    "multimae": "MultiMAE",
    "dino": "DINO",
    "mae": "MAE",
    "vit": "ViT",
}
MODELS_ORDER = ["ViT", "MAE", "DINO", "MultiMAE", "Our-ViT", "Our-MM"]

SCALE_SUFFIXES = {
    "sensormae_base_vit": "Our-ViT-Base",
    "sensormae_small_vit": "Our-ViT-Small",
    "sensormae_tiny_vit": "Our-ViT-Tiny",
}
SCALE_MODELS_ORDER = ["Our-ViT-Base", "Our-ViT-Small", "Our-ViT-Tiny"]

WEATHER_CONDITIONS = ["clear", "light_fog", "dense_fog", "rain/snow"]
WEATHER_LABELS_FULL = ["clear", "light fog", "dense fog", "snow/rain"]
WEATHER_LABELS_SCALE = ["C", "LF", "DF", "S"]
DIFFICULTY_LABELS = ["easy", "moderate", "hard"]


def parse_folder_name(folder_name, model_suffixes):
    """'<n>_det_stf_<modalities>_<backbone>' -> (row label, column label) or (None, None)."""
    match = re.match(r"^\d+_det_stf_(.+)$", folder_name, re.IGNORECASE)
    if not match:
        return None, None

    remainder = match.group(1).lower()
    canonical_model = None
    sensor_str = None
    for suffix in sorted(model_suffixes, key=len, reverse=True):
        marker = "_" + suffix
        if remainder.endswith(marker):
            canonical_model = model_suffixes[suffix]
            sensor_str = remainder[: -len(marker)]
            break

    if canonical_model is None or not sensor_str:
        return None, None

    sensor_key = "+".join(sorted(sensor_str.split("+")))
    canonical_modality = MODALITY_MAP.get(sensor_key)
    return canonical_modality, canonical_model


def load_result(filepath):
    with open(filepath, "r") as f:
        return json.load(f)


def get_car_values(result, weather):
    """Car class easy/moderate/hard AP for a weather condition."""
    if result is None:
        return None, None, None
    weather_data = result.get(weather, {})
    return (
        weather_data.get("Car_image/easy_R40"),
        weather_data.get("Car_image/moderate_R40"),
        weather_data.get("Car_image/hard_R40"),
    )


def get_avg_values(result, weather):
    """All-classes easy/moderate/hard AP for a weather condition."""
    if result is None:
        return None, None, None
    weather_data = result.get(weather, {})
    return (
        weather_data.get("easy_R40"),
        weather_data.get("moderate_R40"),
        weather_data.get("hard_R40"),
    )


def get_avg_weather_value(result, weather):
    """Average easy/moderate/hard AP (all classes) into one number."""
    if result is None:
        return None
    weather_data = result.get(weather, {})
    vals = [weather_data.get(k) for k in ("easy_R40", "moderate_R40", "hard_R40")]
    vals = [float(v) for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


def collect_data(filepaths, model_suffixes):
    """data[modality][model] = results, plus the modalities and models found."""
    data = {}
    discovered_modalities = set()
    discovered_models = set()

    for filepath in filepaths:
        if not os.path.isfile(filepath):
            print(f"File not found, skipping: {filepath}")
            continue

        log_folder = os.path.basename(os.path.dirname(os.path.dirname(filepath)))
        modality, model = parse_folder_name(log_folder, model_suffixes)
        if modality is None or model is None:
            continue

        try:
            result = load_result(filepath)
        except Exception as e:
            print(f"Error loading {filepath}: {e}")
            continue

        data.setdefault(modality, {})
        if model in data[modality]:
            print(f"Warning: duplicate entry for {modality}/{model}, overwriting with {filepath}")
        data[modality][model] = result
        discovered_modalities.add(modality)
        discovered_models.add(model)
        print(f"Loaded [{modality}][{model}]: {filepath}")

    return data, discovered_modalities, discovered_models


def apply_header_style(cell, bold=True, bg_color=None, align="center"):
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if bg_color:
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")


def write_per_difficulty_table(ws, data, modalities, models, value_fn, start_row, title):
    """--mode full: easy / moderate / hard columns per weather condition."""
    n_weather = len(WEATHER_CONDITIONS)

    ws.cell(row=start_row, column=1, value=title)
    apply_header_style(ws.cell(row=start_row, column=1), bold=True, bg_color="DDEBF7", align="left")
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2 + n_weather * 3)
    start_row += 1

    header_row1 = start_row
    ws.cell(row=header_row1, column=1, value="Modalities")
    apply_header_style(ws.cell(row=header_row1, column=1), bg_color="BDD7EE")
    ws.cell(row=header_row1, column=2, value="Model")
    apply_header_style(ws.cell(row=header_row1, column=2), bg_color="BDD7EE")

    col = 3
    for label in WEATHER_LABELS_FULL:
        ws.cell(row=header_row1, column=col, value=label)
        apply_header_style(ws.cell(row=header_row1, column=col), bg_color="BDD7EE")
        ws.merge_cells(start_row=header_row1, start_column=col, end_row=header_row1, end_column=col + 2)
        col += 3

    ws.merge_cells(start_row=header_row1, start_column=1, end_row=header_row1 + 1, end_column=1)
    ws.merge_cells(start_row=header_row1, start_column=2, end_row=header_row1 + 1, end_column=2)

    header_row2 = start_row + 1
    col = 3
    for _ in WEATHER_LABELS_FULL:
        for diff_label in DIFFICULTY_LABELS:
            ws.cell(row=header_row2, column=col, value=diff_label)
            apply_header_style(ws.cell(row=header_row2, column=col), bg_color="DEEAF1")
            col += 1

    current_row = start_row + 2
    for modality in modalities:
        modality_start_row = current_row
        for m_idx, model in enumerate(models):
            if m_idx == 0:
                ws.cell(row=current_row, column=1, value=modality)
                apply_header_style(ws.cell(row=current_row, column=1), bold=False, align="left")

            ws.cell(row=current_row, column=2, value=model)
            ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="left")

            col = 3
            result = data.get(modality, {}).get(model)
            for weather in WEATHER_CONDITIONS:
                easy, moderate, hard = value_fn(result, weather)
                for val in (easy, moderate, hard):
                    cell = ws.cell(row=current_row, column=col, value=float(val) if val is not None else "")
                    cell.alignment = Alignment(horizontal="center")
                    col += 1
            current_row += 1

        if len(models) > 1:
            ws.merge_cells(start_row=modality_start_row, start_column=1, end_row=current_row - 1, end_column=1)
        ws.cell(row=modality_start_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    for i in range(3, 3 + n_weather * 3):
        ws.column_dimensions[get_column_letter(i)].width = 9

    return current_row + 1


def write_averaged_table(ws, data, modalities, models, start_row, title):
    """--mode scale-comparison: one averaged AP per weather condition."""
    n_weather = len(WEATHER_CONDITIONS)

    ws.cell(row=start_row, column=1, value=title)
    apply_header_style(ws.cell(row=start_row, column=1), bold=True, bg_color="DDEBF7", align="left")
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2 + n_weather)
    start_row += 1

    header_row = start_row
    ws.cell(row=header_row, column=1, value="Modalities")
    apply_header_style(ws.cell(row=header_row, column=1), bg_color="BDD7EE")
    ws.cell(row=header_row, column=2, value="Model")
    apply_header_style(ws.cell(row=header_row, column=2), bg_color="BDD7EE")

    col = 3
    for label in WEATHER_LABELS_SCALE:
        ws.cell(row=header_row, column=col, value=label)
        apply_header_style(ws.cell(row=header_row, column=col), bg_color="BDD7EE")
        col += 1

    current_row = header_row + 1
    for modality in modalities:
        modality_start_row = current_row
        for m_idx, model in enumerate(models):
            if m_idx == 0:
                ws.cell(row=current_row, column=1, value=modality)
                apply_header_style(ws.cell(row=current_row, column=1), bold=False, align="left")

            ws.cell(row=current_row, column=2, value=model)
            ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="left")

            col = 3
            result = data.get(modality, {}).get(model)
            for weather in WEATHER_CONDITIONS:
                val = get_avg_weather_value(result, weather)
                cell = ws.cell(row=current_row, column=col, value=round(val, 2) if val is not None else "")
                cell.alignment = Alignment(horizontal="center")
                col += 1
            current_row += 1

        if len(models) > 1:
            ws.merge_cells(start_row=modality_start_row, start_column=1, end_row=current_row - 1, end_column=1)
        ws.cell(row=modality_start_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    for i in range(3, 3 + n_weather):
        ws.column_dimensions[get_column_letter(i)].width = 9

    return current_row + 1


def find_result_files(logs_dir):
    result_files = []
    for folder in sorted(glob.glob(os.path.join(logs_dir, "*_det_stf_*"))):
        candidates = glob.glob(os.path.join(folder, "test_results", "*.json"))
        result_path = min(candidates, key=os.path.getctime) if candidates else None
        if result_path and os.path.isfile(result_path):
            result_files.append(result_path)
        else:
            print(f"Warning: no results file found, skipping folder: {folder}")
    return result_files


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--logs-dir", type=str, default="logs", help="Folder containing STF detection log dirs.")
    parser.add_argument(
        "--mode",
        choices=["full", "scale-comparison"],
        default="full",
        help="'full': all models, per-difficulty AP. 'scale-comparison': SensorMAE-ViT Base/Small/Tiny, averaged AP.",
    )
    parser.add_argument("--output", type=str, default=None, help="Output .xlsx path (default depends on --mode).")
    args = parser.parse_args()

    result_files = find_result_files(args.logs_dir)
    print(f"Found {len(result_files)} result file(s).")

    model_suffixes = MODEL_MAP if args.mode == "full" else SCALE_SUFFIXES
    models_order = MODELS_ORDER if args.mode == "full" else SCALE_MODELS_ORDER

    data, discovered_modalities, discovered_models = collect_data(result_files, model_suffixes)
    if not data:
        print("No matching results found. Exiting.")
        return

    modalities = [m for m in MODALITIES_ORDER if m in discovered_modalities]
    modalities += sorted(discovered_modalities - set(modalities))
    models = [m for m in models_order if m in discovered_models]
    models += sorted(discovered_models - set(models))

    print(f"\nModalities found : {modalities}")
    print(f"Models found     : {models}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    if args.mode == "full":
        next_row = write_per_difficulty_table(
            ws, data, modalities, models, get_car_values, start_row=1, title="Car Class Performance (AP R40)"
        )
        write_per_difficulty_table(
            ws, data, modalities, models, get_avg_values,
            start_row=next_row + 1, title="Average Performance - All Classes (AP R40)",
        )
        default_output = os.path.join(args.logs_dir, "stf_results_summary.xlsx")
    else:
        write_averaged_table(
            ws, data, modalities, models, start_row=1,
            title="Average Performance - All Classes (AP R40, avg easy/moderate/hard)",
        )
        default_output = os.path.join(args.logs_dir, "stf_results_summary_base_small_tiny.xlsx")

    output_path = args.output or default_output
    wb.save(output_path)
    print(f"\nExcel file saved to: {output_path}")


if __name__ == "__main__":
    main()
